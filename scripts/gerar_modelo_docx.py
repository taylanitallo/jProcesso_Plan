# Gera o .docx de exemplo servido em /modelos/termo-referencia.docx para o
# editor OnlyOffice do Modelo Personalizado. Espelha MODELO_TEMPLATE em
# src/mocks/index.ts — atualize os dois juntos se o modelo mudar.
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

BLOCKS = [
    ("h1", "TERMO DE REFERÊNCIA"),
    ("p", "Lei nº 14.133, de 1º de abril de 2021"),
    ("p", "ÓRGÃO OU ENTIDADE PÚBLICA"),
    ("p", "(Processo Administrativo n° xxxxx.xxxxx/xxxx-xx)"),
    ("h2", "1. CONDIÇÕES GERAIS DA CONTRATAÇÃO"),
    ("tag", "objeto"),
    ("tag", "descricao"),
    ("h2", "2. FUNDAMENTAÇÃO E DESCRIÇÃO DA NECESSIDADE DA CONTRATAÇÃO"),
    ("tag", "justificativacontratacao"),
    ("h2", "3. DESCRIÇÃO DA SOLUÇÃO COMO UM TODO CONSIDERADO O CICLO DE VIDA DO OBJETO E ESPECIFICAÇÃO DO PRODUTO"),
    ("tag", "solucaotodo"),
    ("h2", "4. REQUISITOS DA CONTRATAÇÃO"),
    ("tag", "requisitosgarantia"),
    ("h2", "5. MODELO DE EXECUÇÃO DO OBJETO"),
    ("p", "5.1. Condições de Entrega"),
    ("tag", "prazo"),
    ("h2", "6. MODELO DE GESTÃO DO CONTRATO"),
    ("tag", "fiscalizacao"),
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

for kind, texto in BLOCKS:
    if kind == "h1":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.bold = True
        r.font.size = Pt(16)
    elif kind == "h2":
        p = doc.add_paragraph()
        r = p.add_run(texto)
        r.bold = True
        r.font.size = Pt(12)
    elif kind == "tag":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        r = p.add_run(f"[**{texto}**]")
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x8A, 0x6D, 0x1A)
        r.font.highlight_color = 7  # yellow
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(texto)

out_dir = os.path.join(os.path.dirname(__file__), "..", "public", "modelos")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "termo-referencia.docx")
doc.save(out_path)
print("saved:", out_path)
